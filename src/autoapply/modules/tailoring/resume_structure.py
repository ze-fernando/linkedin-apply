import re
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

from autoapply.modules.tailoring.models import (
    Bullet,
    EducationEntry,
    ExperienceEntry,
    ResumeDocument,
    ResumeStructureError,
    SkillLine,
    TextRun,
)
from autoapply.shared.text import normalize_text

SUMMARY = "resumo profissional"
EXPERIENCE = "experiencia profissional"
EDUCATION = "formacao"
SKILLS = "competencias tecnicas"
LANGUAGES = "idiomas"
SECTION_TITLES = {SUMMARY, EXPERIENCE, EDUCATION, SKILLS, LANGUAGES}
BULLET_STYLE = "List Paragraph"
PERIOD_SEPARATOR = "\t"
SKILL_SEPARATOR = re.compile(r",\s*(?![^()]*\))")


def read_runs(paragraph: Paragraph) -> list[TextRun]:
    return [TextRun(text=run.text, bold=bool(run.bold)) for run in paragraph.runs if run.text]


def is_section_title(paragraph: Paragraph) -> bool:
    runs = read_runs(paragraph)
    return len(runs) == 1 and runs[0].bold and PERIOD_SEPARATOR not in runs[0].text


def is_bullet(paragraph: Paragraph) -> bool:
    return paragraph.style is not None and paragraph.style.name == BULLET_STYLE


def split_period(paragraph: Paragraph, kind: str) -> tuple[str, str]:
    title, separator, period = paragraph.text.partition(PERIOD_SEPARATOR)
    if not separator:
        raise ResumeStructureError(f"{kind} sem período: '{paragraph.text}'")
    return title.strip(), period.strip()


def read_bullet(paragraph: Paragraph) -> Bullet:
    runs = read_runs(paragraph)
    if runs and runs[0].bold:
        return Bullet(lead=runs[0].text, text="".join(run.text for run in runs[1:]))
    return Bullet(lead="", text=paragraph.text)


def read_experiences(paragraphs: list[Paragraph]) -> list[ExperienceEntry]:
    entries = []
    for paragraph in paragraphs:
        runs = read_runs(paragraph)
        if is_bullet(paragraph) and entries:
            entries[-1]["bullets"].append(read_bullet(paragraph))
        elif runs and runs[0].bold:
            title, period = split_period(paragraph, "experiência")
            entries.append({"title": title, "period": period, "company": "", "company_detail": "", "bullets": []})
        elif runs and entries:
            entries[-1]["company"] = runs[0].text
            entries[-1]["company_detail"] = "".join(run.text for run in runs[1:])
        else:
            raise ResumeStructureError(f"linha fora de uma experiência: '{paragraph.text}'")
    return [ExperienceEntry(**entry) for entry in entries]


def read_education(paragraphs: list[Paragraph]) -> list[EducationEntry]:
    entries = []
    for paragraph in paragraphs:
        runs = read_runs(paragraph)
        if runs and runs[0].bold:
            course, period = split_period(paragraph, "formação")
            entries.append({"course": course, "period": period, "institution": ""})
        elif entries:
            entries[-1]["institution"] = paragraph.text.strip()
        else:
            raise ResumeStructureError(f"linha fora de uma formação: '{paragraph.text}'")
    return [EducationEntry(**entry) for entry in entries]


def read_skill_line(paragraph: Paragraph) -> SkillLine:
    label, separator, items = paragraph.text.partition(":")
    if not separator:
        raise ResumeStructureError(f"competência sem rótulo: '{paragraph.text}'")
    return SkillLine(label=label.strip(), items=[item.strip() for item in SKILL_SEPARATOR.split(items) if item.strip()])


def split_sections(paragraphs: list[Paragraph]) -> tuple[list[Paragraph], dict[str, list[Paragraph]]]:
    header = []
    sections: dict[str, list[Paragraph]] = {}
    current = None
    for paragraph in paragraphs:
        if is_section_title(paragraph) and (current is not None or header):
            title = normalize_text(paragraph.text.strip())
            if title not in SECTION_TITLES:
                raise ResumeStructureError(f"seção desconhecida: '{paragraph.text.strip()}'")
            current = title
            sections[current] = []
        elif current is None:
            header.append(paragraph)
        else:
            sections[current].append(paragraph)
    return header, sections


def read_resume_document(path: Path) -> ResumeDocument:
    paragraphs = [paragraph for paragraph in docx.Document(str(path)).paragraphs if paragraph.text.strip()]
    header, sections = split_sections(paragraphs)
    missing = [title for title in SECTION_TITLES if title not in sections]
    if len(header) < 2 or missing:
        missing_text = ", ".join(sorted(missing)) or "cabeçalho"
        raise ResumeStructureError(f"currículo fora do layout esperado; faltando: {missing_text}")
    return ResumeDocument(
        name=header[0].text.strip(),
        headline=header[1].text.strip(),
        contact_lines=[paragraph.text.strip() for paragraph in header[2:]],
        summary="\n".join(paragraph.text.strip() for paragraph in sections[SUMMARY]),
        experiences=read_experiences(sections[EXPERIENCE]),
        education=read_education(sections[EDUCATION]),
        skill_lines=[read_skill_line(paragraph) for paragraph in sections[SKILLS]],
        languages=[run for paragraph in sections[LANGUAGES] for run in read_runs(paragraph)],
    )
