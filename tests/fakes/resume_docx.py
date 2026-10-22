from pathlib import Path

import docx
from docx.document import Document

LIST_STYLE = "List Paragraph"


def add_runs(document: Document, runs: list[tuple[str, bool]], style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    for text, bold in runs:
        paragraph.add_run(text).bold = bold or None


def add_heading(document: Document, title: str) -> None:
    add_runs(document, [(title, True)])


def build_resume_docx(path: Path, extra_sections: list[str] | None = None, period_separator: str = "\t") -> Path:
    document = docx.Document()
    if LIST_STYLE not in [style.name for style in document.styles]:
        document.styles.add_style(LIST_STYLE, 1)
    add_runs(document, [("Maria Teste da Silva", True)])
    add_runs(document, [("Desenvolvedora Back-end Python  •  APIs REST  •  Automação", False)])
    add_runs(document, [("São Paulo – SP (remoto)  |  (11) 91234-5678  |  maria.teste@example.com", False)])
    add_runs(document, [("linkedin.com/in/maria-teste  |  github.com/maria-teste", False)])
    add_runs(document, [])
    add_heading(document, "Resumo Profissional")
    add_runs(document, [("Desenvolvedora back-end com foco em Python e APIs REST, com 3 anos de experiência.", False)])
    add_heading(document, "Experiência Profissional")
    add_runs(document, [("Desenvolvedora Pleno", True), (f"{period_separator}Jan 2025 – Atual", False)])
    add_runs(document, [("Empresa Alfa", False), ("  •  fintech", False)])
    add_runs(document, [("Desenvolvo APIs", True), (" em FastAPI com PostgreSQL.", False)], LIST_STYLE)
    add_runs(document, [("Automatizei relatórios", True), (" com Python, reduzindo 2 horas por dia.", False)], LIST_STYLE)
    add_runs(document, [("Desenvolvedora Júnior", True), ("\tMar 2023 – Dez 2024", False)])
    add_runs(document, [("Empresa Beta", False), ("  •  software house", False)])
    add_runs(document, [("Mantive sistemas", True), (" em Django e React.", False)], LIST_STYLE)
    for title in extra_sections or []:
        add_heading(document, title)
    add_heading(document, "Formação")
    add_runs(document, [("Bacharelado em Ciência da Computação", True), ("\tConclusão prevista: Dez 2027", False)])
    add_runs(document, [("Universidade Teste  —  em andamento", False)])
    add_heading(document, "Competências Técnicas")
    add_runs(document, [("Linguagens: ", True), ("Python, SQL, JavaScript", False)])
    add_runs(document, [("Automação & IA: ", True), ("n8n, LLMs (Claude, OpenAI), Selenium", False)])
    add_heading(document, "Idiomas")
    add_runs(document, [("Português: ", True), ("nativo   •   ", False), ("Inglês: ", True), ("intermediário", False)])
    document.save(path)
    return path
