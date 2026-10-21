from collections.abc import Callable
from pathlib import Path

import docx
import pytest
from pypdf import PdfWriter


def escape_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def assemble_pdf(objects: list[str]) -> bytes:
    output = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_offset = len(output)
    xref_entries = "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n{xref_entries}".encode("latin-1")
    output += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("latin-1")
    return output


@pytest.fixture
def build_docx(tmp_path: Path) -> Callable[..., Path]:
    def build(name: str, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> Path:
        document = docx.Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if table_rows:
            table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            for row_cells, row_values in zip(table.rows, table_rows):
                for cell, value in zip(row_cells.cells, row_values):
                    cell.text = value
        path = tmp_path / name
        document.save(path)
        return path

    return build


@pytest.fixture
def build_text_pdf(tmp_path: Path) -> Callable[[str, list[str]], Path]:
    def build(name: str, lines: list[str]) -> Path:
        text_operations = " ".join(f"({escape_pdf_text(line)}) Tj T*" for line in lines)
        content = f"BT /F1 12 Tf 72 720 Td 14 TL {text_operations} ET"
        path = tmp_path / name
        path.write_bytes(
            assemble_pdf(
                [
                    "<< /Type /Catalog /Pages 2 0 R >>",
                    "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                    "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
                    "/Resources << /Font << /F1 5 0 R >> >> >>",
                    f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
                    "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
                ]
            )
        )
        return path

    return build


@pytest.fixture
def build_blank_pdf(tmp_path: Path) -> Callable[[str], Path]:
    def build(name: str) -> Path:
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        path = tmp_path / name
        with path.open("wb") as output:
            writer.write(output)
        return path

    return build
